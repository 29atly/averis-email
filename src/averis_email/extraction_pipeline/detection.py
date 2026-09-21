"""Validate file contents and build explainable, per-page extraction plans."""
from hashlib import file_digest
from pathlib import Path
import posixpath
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile, is_zipfile

import pymupdf

from .models import FileType, PdfPage, PdfType, RoutingPlan

SCAN_IMAGE_FRACTION = 0.50
MIN_TEXT_WORDS_ON_IMAGE_PAGE = 20
MIN_UNCOVERED_IMAGE_FRACTION = 0.15
MAX_BAD_CHARACTER_FRACTION = 0.02
EXTENSIONS = {".pdf": "pdf", ".xlsx": "xlsx", ".txt": "txt", ".docx": "docx", ".docs": "docx"}


def fingerprint(path: Path) -> str:
    with path.open("rb") as stream:
        return file_digest(stream, "sha256").hexdigest()


def decode_text(raw: bytes) -> str:
    """Accept UTF-8 and BOM-marked UTF-16/32; never silently replace bytes."""
    encoding = ("utf-32" if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")) else
                "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
    text = raw.decode(encoding)
    if any(unicodedata.category(c) == "Cc" and c not in "\n\r\t\f" for c in text):
        raise ValueError("Text contains binary control characters")
    return text


def _office_type(path: Path) -> FileType:
    """Identify the package's declared main part, excluding macro formats."""
    with ZipFile(path) as archive:
        def xml(name):
            if archive.getinfo(name).file_size > 1024 * 1024:
                raise ValueError("Office package metadata is too large")
            return ET.fromstring(archive.read(name))
        try:
            relationships = xml("_rels/.rels")
            types = xml("[Content_Types].xml")
        except KeyError:
            return "unsupported"
        mains = [r for r in relationships if r.get("Type", "").endswith("/officeDocument")]
        if len(mains) != 1 or mains[0].get("TargetMode") == "External":
            return "unsupported"
        target = posixpath.normpath(mains[0].get("Target", "").lstrip("/"))
        if target.startswith("../") or target not in archive.namelist():
            raise ValueError("Office package main part is missing or invalid")
        content_types = {r.get("PartName"): r.get("ContentType") for r in types}
        return {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml": "docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml": "xlsx",
        }.get(content_types.get("/" + target), "unsupported")


def recognize_file_type(path: Path) -> FileType:
    """Identify strong binary formats before considering a text filename hint."""
    with path.open("rb") as stream:
        header = stream.read(1024)
    if b"%PDF-" in header:
        return "pdf"
    if is_zipfile(path):
        return _office_type(path)
    if path.suffix.lower() == ".txt" or not path.suffix:
        decode_text(path.read_bytes())
        return "txt"
    if path.suffix.lower() in EXTENSIONS:
        raise ValueError("Contents do not match a supported document format")
    return "unsupported"


def union_area(rectangles) -> float:
    """Rectangle union area: overlapping images/text must not count twice."""
    rects = [r for r in rectangles if not r.is_empty and not r.is_infinite]
    xs = sorted({x for r in rects for x in (r.x0, r.x1)})
    area = 0.0
    for left, right in zip(xs, xs[1:]):
        intervals = sorted((r.y0, r.y1) for r in rects if r.x0 < right and r.x1 > left)
        end = float("-inf")
        height = 0.0
        for low, high in intervals:
            height += max(0, high - max(low, end))
            end = max(end, high)
        area += (right - left) * height
    return area


def text_quality(text: str) -> tuple[bool, float]:
    chars = [c for c in text if not c.isspace()]
    if not chars:
        return False, 0.0
    bad = sum(c == "\ufffd" or unicodedata.category(c) in {"Cc", "Cs", "Co"} for c in chars)
    fraction = bad / len(chars)
    return fraction <= MAX_BAD_CHARACTER_FRACTION, fraction


def inspect_pdf(path: Path) -> tuple[PdfType, list[PdfPage]]:
    pages = []
    with pymupdf.open(path, filetype="pdf") as document:
        if not document.is_pdf:
            raise ValueError("File contents are not a PDF")
        if document.needs_pass:
            raise PermissionError("PDF requires a password")
        for number, page in enumerate(document, 1):
            words = [w for w in page.get_text("words") if w[4].strip()]
            usable, bad = text_quality(" ".join(w[4] for w in words))
            rects = [pymupdf.Rect(w[:4]) & page.rect for w in words]
            images = [pymupdf.Rect(i["bbox"]) & page.rect for i in page.get_image_info()]
            area = page.rect.get_area() or 1
            coverage = union_area(images) / area
            largest = max((r.get_area() / area for r in images), default=0)
            # A full scan with only a header is not a complete text layer.
            sparse_scan = coverage >= SCAN_IMAGE_FRACTION and len(words) < MIN_TEXT_WORDS_ON_IMAGE_PAGE
            uncovered = False
            for image in images:
                if image.get_area() / area < MIN_UNCOVERED_IMAGE_FRACTION:
                    continue
                overlaps = [r & image for r in rects if not (r & image).is_empty]
                span = (max(r.y1 for r in overlaps) - min(r.y0 for r in overlaps)) if overlaps else 0
                if span < image.height * 0.20:
                    uncovered = True
            if words and usable and not sparse_scan and not uncovered:
                kind, method, reason = "text", "native", "usable_embedded_text"
            elif words or images or page.get_drawings():
                kind, method = "scanned", "ocr"
                reason = ("unusable_embedded_text" if words and not usable else
                          "sparse_text_over_scan" if sparse_scan else
                          "image_without_sufficient_text_layer" if uncovered else "visible_content_without_text")
            else:
                # Confirm apparent blanks visually; annotations can be visible
                # even when text/image/drawing inspection finds nothing.
                pix = page.get_pixmap(dpi=36, colorspace=pymupdf.csGRAY, alpha=False)
                visible = any(value < 245 for value in pix.samples)
                kind, method, reason = (("scanned", "ocr", "rendered_content_without_text") if visible else
                                        ("blank", "blank", "blank_render"))
            pages.append(PdfPage(page=number, kind=kind, method=method, word_count=len(words),
                                 largest_image_fraction=round(largest, 4), image_fraction=round(coverage, 4),
                                 text_fraction=round(union_area(rects) / area, 4),
                                 bad_character_fraction=round(bad, 4), reasons=[reason]))
    kinds = {page.kind for page in pages} - {"blank"}
    pdf_type = ("mixed" if kinds == {"text", "scanned"} else
                "epdf" if kinds == {"text"} else "scanned" if kinds == {"scanned"} else "empty")
    return pdf_type, pages


def inspect_document(path: str | Path) -> RoutingPlan:
    path = Path(path).resolve()
    plan = RoutingPlan(source=str(path), supplied_extension=path.suffix.lower())
    try:
        if not path.is_file():
            raise ValueError("Input must be an existing file.")
        plan.sha256 = fingerprint(path)
        plan.file_type = recognize_file_type(path)
        if plan.file_type == "unsupported":
            plan.validation = "unsupported"
            plan.message = "Unsupported document format. Expected PDF, XLSX, DOCX/OOXML, or TXT."
            return plan
        plan.extension_mismatch = EXTENSIONS.get(plan.supplied_extension) != plan.file_type
        if plan.extension_mismatch:
            plan.reasons.append("filename_differs_from_detected_format")
        if plan.file_type == "pdf":
            plan.pdf_type, plan.pages = inspect_pdf(path)
            plan.route = "pdf_extractor" if plan.pdf_type == "epdf" else "ocr"
            if plan.pdf_type == "empty":
                plan.route = None
                plan.validation = "needs_review"
                plan.message = "PDF contains no detected text or visible page content."
                return plan
        elif plan.file_type == "docx":
            from docx import Document
            with path.open("rb") as stream:
                Document(stream)
            plan.route = "docx"
        elif plan.file_type == "xlsx":
            from openpyxl import load_workbook
            with path.open("rb") as stream:
                workbook = load_workbook(stream, read_only=True, data_only=True)
                workbook.close()
            plan.route = "xlsx"
        else:
            decode_text(path.read_bytes())
            plan.route = "txt"
        plan.validation = "valid"
        plan.reasons.append("contents_validated")
    except PermissionError as exc:
        plan.validation = "locked" if plan.file_type == "pdf" and "password" in str(exc) else "invalid"
        plan.message = str(exc)
    except Exception as exc:
        plan.validation = "invalid"
        plan.message = f"Document validation failed: {exc}"
    return plan
