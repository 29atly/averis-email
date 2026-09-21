import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from averis_email.extraction_pipeline import extract_file


class StructuredExtractionTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)

    def test_text_boundaries_multiline_and_duplicates(self):
        path = self.root / 'input.TXT'
        path.write_text('Heading\nShipper (Principal or Seller): Example Ltd\n'
                        '  123 Main St\n\nSingapore\nConsignee:\nNotify Party: Receiver\n'
                        'Commodity: goods with Shipper in the description\n'
                        'Shipper: Second Ltd\nGross Weight: 0 KG\n', encoding='utf-8-sig')
        result = extract_file(path)
        self.assertEqual(result.status, 'EXTRACTED')
        fields = result.extraction['fields']
        self.assertEqual(fields['shipper'], 'Example Ltd\n123 Main St\n\nSingapore')
        self.assertIsNone(fields['consignee'])
        self.assertEqual(fields['notify_party'], 'Receiver')
        self.assertEqual(fields['goods_description'], 'goods with Shipper in the description')
        self.assertEqual(fields['gross_weight_kg'], '0 KG')
        self.assertEqual(len([h for h in result.extraction['occurrences'] if h['field']=='shipper']), 2)
        self.assertNotIn('keyword', result.extraction['occurrences'][0])

    def test_utf16(self):
        path = self.root / 'input.txt'
        path.write_text('Shipper： 公司', encoding='utf-16')
        self.assertEqual(extract_file(path).extraction['fields']['shipper'], '公司')

    def test_word_tables_paragraphs_and_translations_in_order(self):
        doc = Document()
        doc.add_paragraph('B/L NO.(提单号): ABC123')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Shipper (Principal or Seller) (发货人)'
        table.cell(0, 1).text = 'Company\nAddress'
        table.cell(1, 0).text = 'Notify Party/Intermediate Consignee (通知人)'
        table.cell(1, 1).text = 'Receiver'
        doc.add_paragraph('Gross Wt (kgs) (毛重 KGS): 123 KG')
        for extension in ('docx', 'docs'):
            path = self.root / ('input.' + extension)
            doc.save(path)
            result = extract_file(path)
            self.assertEqual(result.status, 'EXTRACTED', result.message)
            fields = result.extraction['fields']
            self.assertEqual(fields['bl_number'], 'ABC123')
            self.assertEqual(fields['shipper'], 'Company\nAddress')
            self.assertEqual(fields['notify_party'], 'Receiver')
            self.assertEqual(fields['gross_weight_kg'], '123 KG')

    def test_spreadsheet_adjacent_cells_exact_labels_and_sheets(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['Heading', 'Ignored'])
        sheet.append(['Shipper:', 'Company\nAddress', 'Unrelated'])
        sheet.append(['Consignee', None, 'Do not skip blank value'])
        sheet.append(['Notify Party', 'Shipper', 'Do not parse value as label'])
        sheet.append(['Gross Weight', 0])
        sheet.append(['Shipper Ltd', 'Not a label'])
        other = book.create_sheet('Second')
        other.append(['Consignee', 'Receiver'])
        other.append(['Shipper', 'Second company'])
        path = self.root / 'input.XLSX'
        book.save(path)
        book.close()
        result = extract_file(path)
        self.assertEqual(result.status, 'EXTRACTED', result.message)
        fields = result.extraction['fields']
        self.assertEqual(fields['shipper'], 'Company\nAddress')
        self.assertEqual(fields['consignee'], 'Receiver')
        self.assertEqual(fields['notify_party'], 'Shipper')
        self.assertEqual(fields['gross_weight_kg'], '0')
        self.assertEqual(len(result.extraction['occurrences']), 6)
        self.assertEqual(result.extraction['occurrences'][-1]['sheet'], 'Second')

    def test_corrupt_office_files(self):
        for extension in ('xlsx', 'docx', 'docs'):
            path = self.root / ('input.' + extension)
            path.write_text('not an Office document')
            self.assertEqual(extract_file(path).status, 'ERROR')
