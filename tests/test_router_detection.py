import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from docx import Document
from openpyxl import Workbook

from averis_email.extraction_pipeline.detection import inspect_document, text_quality, union_area


class RouterDetectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'document.pdf'

    def pdf(self, kind='native'):
        with pymupdf.open() as doc:
            page = doc.new_page(width=600, height=800)
            if kind == 'native':
                page.insert_text((20, 40), 'Shipper: ACME')
            elif kind in ('scan', 'tiles', 'header', 'searchable'):
                with pymupdf.open() as raster:
                    raster.new_page(width=600, height=800).insert_text((20, 40), 'Scanned content')
                    image = raster[0].get_pixmap().tobytes('png')
                if kind == 'tiles':
                    for index in range(4):
                        page.insert_image(pymupdf.Rect(0, index * 200, 600, (index + 1) * 200), stream=image, keep_proportion=False)
                else:
                    page.insert_image(page.rect, stream=image)
                if kind == 'header':
                    page.insert_textbox(pymupdf.Rect(20, 20, 580, 100), 'Header ' * 40, fontsize=8)
                if kind == 'searchable':
                    for y in range(50, 750, 50):
                        page.insert_text((20, y), 'Searchable scanned content on this line', render_mode=3)
            elif kind == 'vector':
                page.draw_rect(pymupdf.Rect(20, 20, 100, 100))
            doc.save(self.path)

    def test_renamed_native_pdf(self):
        self.pdf()
        renamed = self.path.with_suffix('.bin')
        self.path.rename(renamed)
        plan = inspect_document(renamed)
        self.assertEqual(plan.validation, 'valid')
        self.assertEqual(plan.route, 'pdf_extractor')
        self.assertTrue(plan.extension_mismatch)
        self.assertEqual(plan.pages[0].method, 'native')

    def test_pdf_decisions(self):
        for kind, method in [('native', 'native'), ('scan', 'ocr'), ('tiles', 'ocr'),
                             ('header', 'ocr'), ('searchable', 'native'), ('blank', 'blank'), ('vector', 'ocr')]:
            with self.subTest(kind=kind):
                self.pdf(kind)
                plan = inspect_document(self.path)
                self.assertEqual(plan.pages[0].method, method)
                if kind == 'tiles':
                    self.assertEqual(plan.pages[0].image_fraction, 1)
                    self.assertLess(plan.pages[0].largest_image_fraction, .5)
                if kind == 'header':
                    self.assertGreater(plan.pages[0].word_count, 20)

    def test_renamed_office_documents(self):
        doc = Document()
        doc.add_paragraph('Shipper: Example')
        doc.save(self.path)
        self.assertEqual(inspect_document(self.path).file_type, 'docx')
        book = Workbook()
        book.active.append(['Shipper:', 'Example'])
        book.save(self.path)
        plan = inspect_document(self.path)
        self.assertEqual((plan.file_type, plan.validation), ('xlsx', 'valid'))

    def test_invalid_and_unsupported(self):
        self.assertEqual(inspect_document(self.path).validation, 'invalid')
        self.path.write_text('not a PDF')
        self.assertEqual(inspect_document(self.path).validation, 'invalid')
        with ZipFile(self.path, 'w') as archive:
            archive.writestr('random.txt', 'ordinary zip')
        self.assertEqual(inspect_document(self.path).validation, 'unsupported')

    def test_locked(self):
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(self.path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
        self.assertEqual(inspect_document(self.path).validation, 'locked')

    def test_text_encodings_and_binary(self):
        path = self.path.with_suffix('.txt')
        for encoding in ('utf-8', 'utf-16', 'utf-32'):
            path.write_bytes('Shipper: 公司'.encode(encoding))
            self.assertEqual(inspect_document(path).validation, 'valid')
        path.write_bytes(b'abc\x00def')
        self.assertEqual(inspect_document(path).validation, 'invalid')

    def test_multilingual_quality_and_overlapping_coverage(self):
        self.assertTrue(text_quality('公司 日本語 العربية 123')[0])
        self.assertFalse(text_quality('bad \ufffd\ufffd')[0])
        self.assertEqual(union_area([pymupdf.Rect(0, 0, 10, 10), pymupdf.Rect(5, 0, 15, 10)]), 150)


if __name__ == '__main__':
    unittest.main()
