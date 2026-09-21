"""Keyword extraction using logical lines and adjacent cells, without geometry."""
from pathlib import Path
import re

from averis_email.config.keywords import KEYWORD_ALIASES
from averis_email.stages.pdf_extraction import build_extraction_result
from .detection import decode_text

# Additional labels observed in the text and Office samples. Non-focus labels
# also delimit values so trailing freight/order details do not leak into them.
ALIASES = {
    **KEYWORD_ALIASES,
    "notify_party": (*KEYWORD_ALIASES["notify_party"], "Notify Party/Intermediate Consignee"),
    "voyage": (*KEYWORD_ALIASES["voyage"], "Voyage No."),
    "vessel": (*KEYWORD_ALIASES["vessel"], "Export Carrier (vessel, voyage)"),
    "port_of_loading": (*KEYWORD_ALIASES["port_of_loading"], "Port of Loading (POL)"),
    "port_of_discharge": (*KEYWORD_ALIASES["port_of_discharge"], "Port of Discharge (POD)"),
    "freight": ("Freight",),
    "order_number": ("Order No.",),
}


def _patterns():
    for field, variants in ALIASES.items():
        for alias in variants:
            label = r"\s+".join(re.escape(part) for part in alias.split())
            # Samples append Chinese translations, sometimes without a space.
            translation = r"(?:\s*\([^)]*[\u3400-\u9fff][^)]*\))?"
            yield field, alias, re.compile(r"^\s*" + label + translation + r"(?=\s|[:：]|$)", re.I)


PATTERNS = sorted(_patterns(), key=lambda item: -len(item[1]))


def _match(line):
    for field, alias, pattern in PATTERNS:
        match = pattern.match(line)
        if match:
            return field, alias, line[match.end():].lstrip(" \t:：")
    return None


def extract_lines(lines):
    """A label at the start of a line opens a value until the next label."""
    hits = []
    current = None
    parts = []

    def finish():
        if current is not None:
            current["value"] = "\n".join(parts).strip() or None
            hits.append(current)

    for number, line in enumerate(lines, 1):
        match = _match(line)
        if match:
            finish()
            field, alias, value = match
            current = {"field": field, "alias": alias, "line": number}
            parts = [value] if value else []
        elif current is not None:
            parts.append(line.strip())
    finish()
    return hits


def extract_txt(path: Path) -> dict:
    raw = path.read_bytes()
    text = decode_text(raw)
    return build_extraction_result(path, extract_lines(text.splitlines()), [], ALIASES, raw_text=text)


def extract_docx(path: Path) -> dict:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def block_lines(container):
        for block in container.iter_inner_content():
            if isinstance(block, Paragraph):
                yield from block.text.splitlines()
            elif isinstance(block, Table):
                for row in block.rows:
                    seen = set()
                    for cell in row.cells:
                        # A merged cell is repeated by python-docx in a row.
                        if cell._tc not in seen:
                            seen.add(cell._tc)
                            yield from block_lines(cell)

    # Passing a stream also permits OOXML documents named .docs.
    with path.open("rb") as stream:
        document = Document(stream)
        text = "\n".join(block_lines(document))
        hits = extract_lines(text.splitlines())
    return build_extraction_result(path, hits, [], ALIASES, raw_text=text)


def extract_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    stream = path.open("rb")
    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except Exception:
        stream.close()
        raise
    hits = []
    raw_rows = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                raw_rows.append("\t".join(str(cell.value) if cell.value is not None else "" for cell in row))
                index = 0
                while index < len(row):
                    cell = row[index]
                    match = _match(cell.value) if isinstance(cell.value, str) else None
                    if match and not match[2].strip():
                        field, alias, _ = match
                        value = row[index + 1].value if index + 1 < len(row) else None
                        hits.append({"field": field, "alias": alias, "sheet": sheet.title,
                                     "cell": cell.coordinate,
                                     "value": str(value).strip() if value is not None else None})
                        # The adjacent value can itself be a keyword; never
                        # reinterpret that cell as a second label.
                        index += 2
                    else:
                        index += 1
    finally:
        workbook.close()
        stream.close()
    return build_extraction_result(path, hits, [], ALIASES, raw_text="\n".join(raw_rows))
