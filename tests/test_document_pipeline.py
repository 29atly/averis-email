"""Real format adapters and ingestion; inference is replaced at the model boundary."""
from io import BytesIO
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import pymupdf
from docx import Document
from openpyxl import Workbook

from averis_email.extraction_pipeline import extract_file
from averis_email.orchestrator import run_pipeline
from averis_email.stages.ingestion import read_document

TEXT = ('Shipper: Example Ltd\nConsignee: Receiver Ltd\nNotify Party: Notify Ltd\n'
        'Port of Loading: Singapore\nPort of Discharge: Rotterdam\n'
        'No. of Containers: 2\nGross Weight: 1000 KG')


class Loader:
    def __init__(self, files):
        self.files = files

    def read_bytes(self, path):
        return self.files[path]


def office_bytes(kind):
    stream = BytesIO()
    if kind == 'xlsx':
        book = Workbook()
        for line in TEXT.splitlines():
            label, value = line.split(': ', 1)
            book.active.append([label, value])
        book.save(stream)
        book.close()
    else:
        doc = Document()
        for line in TEXT.splitlines():
            doc.add_paragraph(line)
        doc.save(stream)
    return stream.getvalue()


def pdf_bytes(scan=False):
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(20, 20, 550, 750), TEXT)
        if not scan:
            return doc.tobytes()
        image = page.get_pixmap().tobytes('png')
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_image(page.rect, stream=image)
        return doc.tobytes()


class FakeOCR:
    def predict(self, image, **kwargs):
        lines = TEXT.splitlines()
        return [{'rec_texts': lines, 'rec_scores': [0.99] * len(lines),
                 'rec_boxes': [[20, 20+i*50, 600, 45+i*50] for i in range(len(lines))]}]


class DocumentPipelineTests(unittest.TestCase):
    def test_all_document_formats_flow_through_ingestion(self):
        files = {'input.txt': TEXT.encode(), 'input.docx': office_bytes('docx'),
                 'input.docs': office_bytes('docx'), 'input.xlsx': office_bytes('xlsx'),
                 'native.pdf': pdf_bytes(), 'scan.pdf': pdf_bytes(True)}
        with patch('averis_email.extraction_pipeline.ocr.create_ocr_engine', return_value=FakeOCR()):
            for name in files:
                with self.subTest(name=name):
                    doc = read_document(Loader(files), name)
                    self.assertTrue(doc.readable, doc.error)
                    self.assertEqual(doc.attachment_path, name)
                    self.assertIn('Example Ltd', doc.fields['_raw'])
                    self.assertIn('1000 KG', doc.fields['_raw'])

    def test_renamed_workbook_and_utf32_use_content_validated_adapters(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'renamed.bin'
            path.write_bytes(office_bytes('xlsx'))
            result = extract_file(path)
            self.assertEqual(result.status, 'EXTRACTED', result.message)
            self.assertEqual(result.extraction['fields']['shipper'], 'Example Ltd')
            self.assertEqual(result.extraction['occurrences'][0]['cell'], 'A1')
            path = path.with_suffix('.txt')
            path.write_bytes(TEXT.encode('utf-32'))
            self.assertEqual(extract_file(path).extraction['fields']['shipper'], 'Example Ltd')

    def test_comparison_pipeline_uses_router_and_real_field_extraction(self):
        files = {'SI.xlsx': office_bytes('xlsx'), 'BL.docs': office_bytes('docx')}
        with (patch('averis_email.orchestrator.classify_email',
                    return_value=SimpleNamespace(category='BL_COMPARISON', review_required=False)),
              patch('averis_email.stages.extraction._extract_with_gemini', return_value=None)):
            result = run_pipeline(Loader(files), {'email_id': 'routed', 'attachments': list(files)})
        self.assertEqual(result.status, 'OK', result.model_dump())
        self.assertEqual(result.si.fields['shipper'].value, 'Example Ltd')
        self.assertEqual(result.bl.fields['shipper'].source_file, 'BL.docs')

    def test_invalid_blank_and_missing_attachments_are_unreadable(self):
        for name, data in [('bad.docs', b'bad'), ('empty.txt', b'')]:
            doc = read_document(Loader({name: data}), name)
            self.assertFalse(doc.readable)
            self.assertTrue(doc.error)
        self.assertFalse(read_document(Loader({}), 'missing.pdf').readable)

    def test_partial_ocr_cannot_proceed_to_comparison(self):
        with patch('averis_email.extraction_pipeline.ocr._page_words_from_ocr', return_value=[]), \
             patch('averis_email.extraction_pipeline.ocr.create_ocr_engine', return_value=FakeOCR()):
            doc = read_document(Loader({'scan.pdf': pdf_bytes(True)}), 'scan.pdf')
        self.assertFalse(doc.readable)
        self.assertIn('no usable text', doc.error)
